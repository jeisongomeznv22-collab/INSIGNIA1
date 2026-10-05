"""
Lectura del Excel de entrada (IDs a consultar) y escritura del Excel de
reporte final, con deteccion flexible de columnas para poder usar tanto un
Excel simple de una sola columna como el "Reporte de Envios" completo que
exporta el proveedor 4-72.
"""
from __future__ import annotations

import dataclasses
import re
import unicodedata
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_ID_HEADERS = {"id mensaje", "id de correo", "id correo", "id", "codigo", "id del correo",
               # 29-sep (Andres, IDS_PROCESAMIENTO_LOTE_3.xlsx): una sola columna 'IDS'
               "ids", "id's", "id mensajes", "ids mensaje", "ids portal antiguo", "id portal antiguo",
               "id portal nuevo", "ids portal nuevo", "identificador", "id envio", "id del mensaje"}
_DEST_HEADERS = {"destinatario", "correo destinatario", "email"}
_ESTADO_HEADERS = {"estado entrega", "estado", "estado de entrega"}
_ASUNTO_HEADERS = {"asunto"}

# Portal nuevo: hexadecimal de 40. Portal antiguo: numero corto ('1371220').
_ID_PATTERN = re.compile(r"^(?:[0-9A-Fa-f]{20,64}|\d{4,15})$")


def _texto_id(v) -> str:
    """Celda -> ID en texto. Un numero que Excel guardo como 1371220.0 queda
    '1371220' (el portal antiguo usa IDs numericos)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _norm(text) -> str:
    if text is None:
        return ""
    text = str(text).strip().lower()
    text = "".join(
        c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn"
    )
    return text


@dataclasses.dataclass
class InputRow:
    message_id: str
    destinatario: str = ""
    estado_entrega: str = ""
    asunto: str = ""


def read_ids_from_excel(path: str, sheet_name: Optional[str] = None) -> list[InputRow]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]

    rows_iter = ws.iter_rows(values_only=True)
    all_rows = list(rows_iter)
    wb.close()

    if not all_rows:
        return []

    header_row_idx = None
    col_map: dict[str, int] = {}
    for i, row in enumerate(all_rows[:10]):
        candidate = {}
        for j, cell in enumerate(row):
            norm = _norm(cell)
            if norm in _ID_HEADERS:
                candidate["id"] = j
            elif norm in _DEST_HEADERS:
                candidate["destinatario"] = j
            elif norm in _ESTADO_HEADERS:
                candidate["estado"] = j
            elif norm in _ASUNTO_HEADERS:
                candidate["asunto"] = j
        if "id" in candidate:
            header_row_idx = i
            col_map = candidate
            break

    results: list[InputRow] = []
    seen: set[str] = set()

    if header_row_idx is not None:
        data_rows = all_rows[header_row_idx + 1 :]
        id_col = col_map["id"]
        dest_col = col_map.get("destinatario")
        estado_col = col_map.get("estado")
        asunto_col = col_map.get("asunto")
        for row in data_rows:
            if id_col >= len(row):
                continue
            raw_id = row[id_col]
            if raw_id is None:
                continue
            message_id = _texto_id(raw_id)
            if not message_id or message_id in seen:
                continue
            seen.add(message_id)
            results.append(
                InputRow(
                    message_id=message_id,
                    destinatario=str(row[dest_col]).strip() if dest_col is not None and dest_col < len(row) and row[dest_col] else "",
                    estado_entrega=str(row[estado_col]).strip() if estado_col is not None and estado_col < len(row) and row[estado_col] else "",
                    asunto=str(row[asunto_col]).strip() if asunto_col is not None and asunto_col < len(row) and row[asunto_col] else "",
                )
            )
    else:
        # No se reconocio un encabezado: asumir que la columna A trae los IDs
        # (tolerando o no una fila de encabezado simple sin nombre reconocido).
        for row in all_rows:
            if not row:
                continue
            raw_id = row[0]
            if raw_id is None:
                continue
            message_id = _texto_id(raw_id)
            if not message_id or message_id in seen:
                continue
            if not _ID_PATTERN.match(message_id):
                # probablemente la fila de encabezado; se ignora
                continue
            seen.add(message_id)
            results.append(InputRow(message_id=message_id))

    return results


# ---------------------------------------------------------------------------
# Reporte de salida
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_THIN = Side(style="thin", color="C9C9C9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

REPORT_COLUMNS = [
    "Tipo de AA",
    "Numero",
    "Fecha",
    "Nombre Titular",
    "Destinatario",
    "ID Mensaje",
    "Archivo Generado",
    "Estado",
    "Observacion",
]


def write_report_excel(path: str, rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Procesamiento"

    for col_idx, title in enumerate(REPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    for r, row in enumerate(rows, start=2):
        values = [
            row.get("tipo", ""),
            row.get("numero", ""),
            row.get("fecha", ""),
            row.get("nombre", ""),
            row.get("destinatario", ""),
            row.get("id_mensaje", ""),
            row.get("archivo", ""),
            row.get("estado", ""),
            row.get("observacion", ""),
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    widths = [14, 12, 12, 26, 26, 34, 40, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(path)


LOCAL_REPORT_COLUMNS = [
    "Tipo de AA",
    "Numero",
    "Fecha",
    "Nombre Titular",
    "Destinatario",
    "Origen (archivo)",
    "Archivo Generado",
    "Estado",
    "Observacion",
]


def write_local_report_excel(path: str, rows: list[dict]) -> None:
    """Igual que write_report_excel, pero para el lote de ZIPs locales:
    la columna clave es el nombre del archivo/zip de origen en vez del
    ID Mensaje del portal (aqui no hay conexion al portal de por medio)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Procesamiento"

    for col_idx, title in enumerate(LOCAL_REPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    for r, row in enumerate(rows, start=2):
        values = [
            row.get("tipo", ""),
            row.get("numero", ""),
            row.get("fecha", ""),
            row.get("nombre", ""),
            row.get("destinatario", ""),
            row.get("origen", ""),
            row.get("archivo", ""),
            row.get("estado", ""),
            row.get("observacion", ""),
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    widths = [14, 12, 12, 26, 26, 40, 40, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(path)
